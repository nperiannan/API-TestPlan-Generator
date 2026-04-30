"""
Extract sub-feature test cases from a parent YAML and create a standalone test plan.

Usage:
    python extract_subfeature.py <parent.yaml> <subfeature-name> <keyword1> [keyword2 ...]

Examples:
    python extract_subfeature.py Testplans/port.yaml port-poe poe power-over-ethernet power
    python extract_subfeature.py Testplans/port.yaml port-elrp elrp
    python extract_subfeature.py Testplans/port.yaml port-slpp slpp
    python extract_subfeature.py Testplans/port.yaml port-stp stp bpdu enable-edge path-cost
    python extract_subfeature.py Testplans/port.yaml port-mac-locking mac-lock mac-locking

Sub-feature YAML, CSV, and XLSX files are written to Testplans/<subfeature-name>.*
"""
import csv
import json
import os
import sys

import yaml

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

# ── Excel formatting ─────────────────────────────────────────────────────
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=11)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

COLUMNS = [
    "Test Case ID", "Type", "Priority", "Automation", "Description",
    "Method", "Endpoint", "Request Body", "Expected Result",
]
COL_WIDTHS = {"A": 16, "B": 14, "C": 10, "D": 14, "E": 60,
              "F": 10, "G": 50, "H": 60, "I": 60}


def matches_keywords(test, keywords):
    """Check if test description or step body references any keyword."""
    desc = test.get("description", "").lower()
    for kw in keywords:
        if kw in desc:
            return True
    # Also check step body property names
    for step in test.get("steps", []):
        body = step.get("body")
        if isinstance(body, dict):
            for obj in body.get("objects", []):
                if isinstance(obj, dict):
                    for prop in obj.get("properties", []):
                        if isinstance(prop, dict):
                            pname = prop.get("name", "").lower()
                            for kw in keywords:
                                if kw in pname:
                                    return True
    return False


def fmt_body(body):
    if not body:
        return ""
    if isinstance(body, dict):
        return json.dumps(body, indent=2)
    return str(body)


def fmt_validations(steps):
    parts = []
    for step in steps:
        for v in step.get("validations", []):
            if isinstance(v, dict):
                d = v.get("description", "")
                if d:
                    parts.append(d)
                else:
                    parts.append(f"Status {v.get('expected', '')}")
            elif isinstance(v, str):
                parts.append(v)
    return "\n".join(parts) if parts else ""


def extract_row(tc):
    steps = tc.get("steps", [])
    methods, paths, bodies = [], [], []
    for s in steps:
        m = s.get("method", "")
        p = s.get("path", "")
        if m:
            methods.append(m)
        if p:
            paths.append(p)
        b = s.get("body")
        if b:
            bodies.append(fmt_body(b))
    return [
        tc.get("testCaseID", ""),
        tc.get("type", ""),
        tc.get("priority", ""),
        tc.get("automation", "Automatable"),
        tc.get("description", ""),
        " / ".join(dict.fromkeys(methods)),
        "\n".join(dict.fromkeys(paths)),
        "\n---\n".join(bodies) if bodies else "",
        fmt_validations(steps),
    ]


def write_xlsx(by_cat, xlsx_path):
    wb = Workbook()
    wb.remove(wb.active)
    total = 0
    for cat, rows in by_cat.items():
        ws = wb.create_sheet(title=cat.title())
        ws.append(COLUMNS)
        for col_letter, width in COL_WIDTHS.items():
            ws.column_dimensions[col_letter].width = width
        for cell in ws[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER
        for row in rows:
            ws.append(row)
        for r in range(2, ws.max_row + 1):
            for c in range(1, len(COLUMNS) + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = BODY_FONT
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = THIN_BORDER
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        total += len(rows)
    wb.save(xlsx_path)
    return total


def main():
    if len(sys.argv) < 4:
        print(__doc__)
        sys.exit(1)

    parent_yaml = sys.argv[1]
    subfeature = sys.argv[2]
    keywords = [kw.lower() for kw in sys.argv[3:]]

    out_dir = os.path.dirname(parent_yaml) or "Testplans"
    yaml_out = os.path.join(out_dir, f"{subfeature}.yaml")
    csv_out = os.path.join(out_dir, f"{subfeature}.csv")
    xlsx_out = os.path.join(out_dir, f"{subfeature}.xlsx")

    print(f"Reading: {parent_yaml}")
    print(f"Keywords: {keywords}")

    with open(parent_yaml, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)

    feat = data["features"][0]
    cats = feat["tests"]

    # Collect matching test cases
    by_cat = {}
    yaml_tests = {}
    all_rows = []
    for cat, tests in cats.items():
        matched = [t for t in tests if matches_keywords(t, keywords)]
        if matched:
            by_cat[cat] = [extract_row(t) for t in matched]
            yaml_tests[cat] = matched
            all_rows.extend(by_cat[cat])
            print(f"  {cat.title()}: {len(matched)} tests")

    if not all_rows:
        print("No matching test cases found.")
        return

    # Write sub-feature YAML
    sub_yaml = {
        "version": data.get("version", "1.0"),
        "generatedAt": data.get("generatedAt", ""),
        "sourceYangDir": data.get("sourceYangDir", ""),
        "sourceRESTAPI": data.get("sourceRESTAPI", ""),
        "sourceNOSAPI": data.get("sourceNOSAPI", ""),
        "features": [{
            "featureName": subfeature,
            "featurePath": feat.get("featurePath", ""),
            "profileType": feat.get("profileType", ""),
            "blueprintCategory": feat.get("blueprintCategory", ""),
            "description": f"Sub-feature extracted from {feat.get('featureName', '')}",
            "tests": yaml_tests,
        }],
    }
    with open(yaml_out, "w", encoding="utf-8") as f:
        yaml.dump(sub_yaml, f, default_flow_style=False, allow_unicode=True, sort_keys=False, width=200)
    print(f"Wrote {yaml_out}")

    # Write CSV
    with open(csv_out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        w.writerows(all_rows)
    print(f"Wrote {len(all_rows)} rows to {csv_out}")

    # Write XLSX
    total = write_xlsx(by_cat, xlsx_out)
    print(f"Wrote {xlsx_out} ({len(by_cat)} sheets, {total} total test cases)")


if __name__ == "__main__":
    main()
