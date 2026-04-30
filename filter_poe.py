"""Filter PoE test cases from port.yaml and export to CSV and XLSX."""
import csv
import sys
import yaml

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

INPUT = "Testplans/port.yaml"
CSV_OUT = "Testplans/port-poe.csv"
XLSX_OUT = "Testplans/port-poe.xlsx"

POE_KEYWORDS = ["poe", "power-over-ethernet", "power over ethernet", "power"]

COLUMNS = [
    "Test Case ID",
    "Type",
    "Priority",
    "Automation",
    "Description",
    "Method",
    "Endpoint",
    "Request Body",
    "Expected Result",
]

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=11)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

COL_WIDTHS = {
    "A": 16, "B": 14, "C": 10, "D": 14, "E": 60,
    "F": 10, "G": 50, "H": 60, "I": 60,
}


def is_poe(test):
    desc = test.get("description", "").lower()
    return any(kw in desc for kw in POE_KEYWORDS)


def fmt_body(body):
    if not body:
        return ""
    if isinstance(body, dict):
        import json
        return json.dumps(body, indent=2)
    return str(body)


def fmt_validations(steps):
    parts = []
    for step in steps:
        for v in step.get("validations", []):
            if isinstance(v, dict):
                d = v.get("description", "")
                if d:
                    parts.append(d)
                else:
                    parts.append(f"Status {v.get('expected', '')}")
            elif isinstance(v, str):
                parts.append(v)
    return "\n".join(parts) if parts else ""


def extract_row(tc):
    steps = tc.get("steps", [])
    methods = []
    paths = []
    bodies = []
    for s in steps:
        m = s.get("method", "")
        p = s.get("path", "")
        if m:
            methods.append(m)
        if p:
            paths.append(p)
        b = s.get("body")
        if b:
            bodies.append(fmt_body(b))

    return [
        tc.get("testCaseID", ""),
        tc.get("type", ""),
        tc.get("priority", ""),
        tc.get("automation", "Automatable"),
        tc.get("description", ""),
        " / ".join(dict.fromkeys(methods)),  # unique, order-preserved
        "\n".join(dict.fromkeys(paths)),
        "\n---\n".join(bodies) if bodies else "",
        fmt_validations(steps),
    ]


def main():
    print(f"Reading: {INPUT}")
    with open(INPUT, "r") as f:
        data = yaml.safe_load(f)

    feat = data["features"][0]
    cats = feat["tests"]

    # ── Collect PoE rows by category ──
    by_cat = {}
    all_rows = []
    for cat, tests in cats.items():
        poe_tests = [t for t in tests if is_poe(t)]
        if poe_tests:
            rows = [extract_row(t) for t in poe_tests]
            by_cat[cat] = rows
            all_rows.extend(rows)
            print(f"  {cat.title()}: {len(poe_tests)} PoE tests")

    if not all_rows:
        print("No PoE test cases found.")
        return

    # ── CSV ──
    with open(CSV_OUT, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        w.writerows(all_rows)
    print(f"Wrote {len(all_rows)} rows to {CSV_OUT}")

    # ── XLSX ──
    wb = Workbook()
    wb.remove(wb.active)

    for cat, rows in by_cat.items():
        ws = wb.create_sheet(title=cat.title())
        ws.append(COLUMNS)
        for col_letter, width in COL_WIDTHS.items():
            ws.column_dimensions[col_letter].width = width

        # Header style
        for cell in ws[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN_BORDER

        # Data rows
        for row in rows:
            ws.append(row)

        for r in range(2, ws.max_row + 1):
            for c in range(1, len(COLUMNS) + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = BODY_FONT
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = THIN_BORDER

        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"

    wb.save(XLSX_OUT)
    print(f"Wrote {XLSX_OUT} ({len(by_cat)} sheets, {len(all_rows)} total PoE test cases)")


if __name__ == "__main__":
    main()
