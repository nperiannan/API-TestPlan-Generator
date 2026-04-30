"""
Convert radius-server.yaml test cases to Excel format matching the
NVO-11239 template structure.

Columns: Test Case ID, Testcase Title, Status, Type, Description, Precondition,
         Test Step Description, Test Step Expected Result, Priority

Each test category (boundary, functional, negative, performance, scale)
goes into its own sheet.
"""

import yaml
import json
import re
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def format_body_as_payload(body):
    """Format the request body into a readable payload string."""
    if not body:
        return ""
    # Build a clean representation
    lines = []
    feature_path = body.get('featurePath', '')
    object_type = body.get('objectType', '')
    operation = body.get('operation', '')
    objects = body.get('objects', [])

    payload = {}
    if feature_path:
        payload['featurePath'] = feature_path
    if object_type:
        payload['objectType'] = object_type
    if operation:
        payload['operation'] = operation
    if objects:
        clean_objects = []
        for obj in objects:
            clean_obj = {}
            if obj.get('operation'):
                clean_obj['operation'] = obj['operation']
            if obj.get('type'):
                clean_obj['type'] = obj['type']
            props = obj.get('properties', [])
            if props:
                clean_obj['properties'] = []
                for p in props:
                    prop_entry = {
                        'name': p.get('name', ''),
                        'value': p.get('value', '')
                    }
                    clean_obj['properties'].append(prop_entry)
            clean_objects.append(clean_obj)
        payload['objects'] = clean_objects

    return json.dumps(payload, indent=2)


def format_step_description(steps):
    """Convert YAML steps into human-readable test step descriptions."""
    lines = []
    for i, step in enumerate(steps, 1):
        method = step.get('method', 'GET')
        path = step.get('path', '')
        desc = step.get('description', '')
        body = step.get('body', {})
        expected_status = step.get('expectedStatus', '')
        timeout = step.get('timeout', None)

        lines.append(f"{i}) {desc}")
        lines.append(f"   {method} {{base_url}}{path}")

        if body:
            payload_str = format_body_as_payload(body)
            lines.append(f"   Payload: {payload_str}")

        if expected_status:
            lines.append(f"   Expected HTTP Status: {expected_status}")

        if timeout:
            lines.append(f"   Timeout: {timeout}ms")

        lines.append("")

    return "\n".join(lines).strip()


def format_expected_results(steps):
    """Extract validations (excluding assertions) as expected results."""
    results = []
    for i, step in enumerate(steps, 1):
        validations = step.get('validations', [])
        if validations:
            if len(steps) > 1:
                results.append(f"Step {i}:")
            for v in validations:
                results.append(f"- {v}")
            if len(steps) > 1:
                results.append("")

    return "\n".join(results).strip()


def apply_header_style(ws):
    """Apply styling to the header row."""
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col in range(1, 10):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border


def apply_data_style(ws, max_row):
    """Apply styling to data cells."""
    wrap_alignment = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for row in range(2, max_row + 1):
        for col in range(1, 10):
            cell = ws.cell(row=row, column=col)
            cell.alignment = wrap_alignment
            cell.border = thin_border


def compress_title(description):
    """Compress verbose descriptions into concise titles."""
    title = description

    # Remove 'Boundary test: ' / 'Negative test: ' prefixes (type is already in the sheet name)
    title = re.sub(r'^(?:Boundary|Negative|Performance|Scale)\s+test:\s*', '', title, flags=re.IGNORECASE)

    # Remove long quoted values — keep just the parameter and label
    # e.g. server='2001:db8:ffff:...:fffe' — Maximum valid IPv6 ... → server — Max valid IPv6 unicast
    def shorten_quoted(m):
        param = m.group(1)
        label = m.group(3) if m.group(3) else ''
        return f"{param} {label}".strip()
    title = re.sub(r"(\w[\w-]*)='[^']{20,}'\s*[—–-]\s*(.*)", r'\1 — \2', title)

    # Shorten common verbose phrases
    title = title.replace('(independent test)', '').strip()
    title = title.replace('(above 0.0.0.0/8 reserved range)', '').strip()
    title = title.replace('(last address before multicast 224/4)', '').strip()
    title = title.replace('(RFC 3849 documentation prefix)', '').strip()
    title = title.replace('within documentation prefix', '').strip()
    title = title.replace('(253 characters per RFC 1035)', '').strip()
    title = title.replace('(single-character label + TLD)', '').strip()
    title = re.sub(r'\s*\(create to maximum allowed\)', '', title)
    title = re.sub(r'\s*\(Create -> Read -> Update -> Read -> Delete\)', '', title)
    title = re.sub(r'\s*\(only update specific fields\)', '', title)

    # Shorten "Minimum valid" → "Min valid", "Maximum valid" → "Max valid" etc.
    title = title.replace('Minimum valid', 'Min valid')
    title = title.replace('Maximum valid', 'Max valid')
    title = title.replace('Minimum-length', 'Min-length')
    title = title.replace('Maximum-length', 'Max-length')
    title = title.replace('minimum', 'min').replace('maximum', 'max')

    # Remove trailing em-dashes with nothing after
    title = re.sub(r'\s*[—–-]\s*$', '', title)

    # Collapse multiple spaces
    title = re.sub(r'\s{2,}', ' ', title).strip()

    return title


def set_column_widths(ws):
    """Set reasonable column widths."""
    widths = {
        1: 16,   # Test Case ID
        2: 50,   # Testcase Title
        3: 18,   # Status
        4: 12,   # Type
        5: 60,   # Description
        6: 50,   # Precondition
        7: 80,   # Test Step Description
        8: 60,   # Test Step Expected Result
        9: 10,   # Priority
    }
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def create_sheet_for_category(wb, category_name, test_cases):
    """Create an Excel sheet for a test category."""
    # Clean sheet name (max 31 chars for Excel)
    sheet_name = category_name.capitalize()
    if len(sheet_name) > 31:
        sheet_name = sheet_name[:31]

    ws = wb.create_sheet(title=sheet_name)

    # Headers
    headers = [
        "Test Case ID", "Testcase Title", "Status", "Type", "Description",
        "Precondition", "Test Step Description",
        "Test Step Expected Result", "Priority"
    ]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)

    apply_header_style(ws)

    # Data rows
    for row_idx, tc in enumerate(test_cases, 2):
        description = tc.get('description', '')
        priority = tc.get('priority', 'P3')
        steps = tc.get('steps', [])
        test_case_id = tc.get('testCaseID', '')

        # Title: compressed description (no test case ID)
        title = compress_title(description)

        # Test Step Description: human-readable format
        step_desc = format_step_description(steps)

        # Expected Result: validations only (no assertions)
        expected_result = format_expected_results(steps)

        ws.cell(row=row_idx, column=1, value=test_case_id)
        ws.cell(row=row_idx, column=2, value=title)
        ws.cell(row=row_idx, column=3, value="To Be Automated")
        ws.cell(row=row_idx, column=4, value="Manual")
        ws.cell(row=row_idx, column=5, value=description)
        ws.cell(row=row_idx, column=6, value="QA environment available with EP1-NGC Framework integration")
        ws.cell(row=row_idx, column=7, value=step_desc)
        ws.cell(row=row_idx, column=8, value=expected_result)
        ws.cell(row=row_idx, column=9, value=priority)

    apply_data_style(ws, len(test_cases) + 1)
    set_column_widths(ws)

    return ws


def main():
    import sys, os
    if len(sys.argv) >= 2:
        yaml_path = sys.argv[1]
    else:
        yaml_path = r'c:\Users\nperiannan\OneDrive - Extreme Networks, Inc\Desktop\Testcase generator\Testplans\radius-server.yaml'

    if len(sys.argv) >= 3:
        output_path = sys.argv[2]
    else:
        output_path = os.path.splitext(yaml_path)[0] + '.xlsx'

    print(f"Reading YAML: {yaml_path}")
    with open(yaml_path, 'r', encoding='utf-8') as f:
        data = yaml.safe_load(f)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    features = data.get('features', [])
    total_tests = 0

    # Define category order
    category_order = ['functional', 'boundary', 'negative', 'performance', 'scale']

    for feat in features:
        feature_name = feat.get('featureName', 'unknown')
        tests = feat.get('tests', {})

        for category in category_order:
            test_list = tests.get(category, [])
            if not test_list:
                continue

            print(f"  Processing {category}: {len(test_list)} test cases")
            create_sheet_for_category(wb, category, test_list)
            total_tests += len(test_list)

        # Handle any categories not in the predefined order
        for category, test_list in tests.items():
            if category not in category_order and test_list:
                print(f"  Processing {category}: {len(test_list)} test cases")
                create_sheet_for_category(wb, category, test_list)
                total_tests += len(test_list)

    try:
        wb.save(output_path)
    except PermissionError:
        # File may be open; save with a suffix
        alt_path = output_path.replace('.xlsx', '_v2.xlsx')
        wb.save(alt_path)
        output_path = alt_path
        print(f"  (original file was locked, saved as {alt_path})")
    print(f"\nDone! Generated {output_path}")
    print(f"Total test cases: {total_tests}")
    print(f"Sheets: {wb.sheetnames}")


if __name__ == '__main__':
    main()
